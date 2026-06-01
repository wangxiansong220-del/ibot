from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DOC_PATH = ROOT / "output" / "doc" / "ibot_thesis_draft.docx"
TARGET_DOC_PATH = ROOT / "output" / "doc" / "ibot_thesis_with_testshots.docx"
SHOT_DIR = ROOT / "output" / "playwright"

CN_FONT = "宋体"
HEI_FONT = "黑体"
EN_FONT = "Times New Roman"


def set_run_font(run, size: float, east_asia: str = CN_FONT, ascii_font: str = EN_FONT, bold: bool = False) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def add_picture_after(paragraph: Paragraph, image_path: Path, width_cm: float = 13.6) -> Paragraph:
    p = insert_paragraph_after(paragraph)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    return p


def add_caption_after(paragraph: Paragraph, caption: str) -> Paragraph:
    p = insert_paragraph_after(paragraph)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(caption)
    set_run_font(run, 10.5, east_asia=HEI_FONT)
    return p


def add_body_after(paragraph: Paragraph, text: str) -> Paragraph:
    p = insert_paragraph_after(paragraph)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = Pt(22)
    run = p.add_run(text)
    set_run_font(run, 12)
    return p


def find_paragraph(doc: Document, text: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip() == text:
            return para
    raise ValueError(f"Could not find paragraph: {text}")


def has_caption(doc: Document, caption: str) -> bool:
    return any(para.text.strip() == caption for para in doc.paragraphs)


def replace_image_before_caption(doc: Document, caption: str, image_path: Path, width_cm: float) -> None:
    paragraphs = list(doc.paragraphs)
    for idx, para in enumerate(paragraphs):
        if para.text.strip() == caption and idx > 0:
            old_image_para = paragraphs[idx - 1]
            new_p = OxmlElement("w:p")
            old_image_para._p.addprevious(new_p)
            new_para = Paragraph(new_p, old_image_para._parent)
            new_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            new_para.paragraph_format.first_line_indent = Cm(0)
            new_para.paragraph_format.left_indent = Cm(0)
            new_para.paragraph_format.space_before = Pt(6)
            new_para.paragraph_format.space_after = Pt(6)
            new_para.paragraph_format.keep_together = True
            new_para.paragraph_format.keep_with_next = True
            new_para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = new_para.add_run()
            run.add_picture(str(image_path), width=Cm(width_cm))
            old_image_para._element.getparent().remove(old_image_para._element)
            return


def insert_figure_block(
    doc: Document,
    anchor_text: str,
    image_name: str,
    caption: str,
    note: str,
    width_cm: float = 13.6,
) -> None:
    if has_caption(doc, caption):
        replace_image_before_caption(doc, caption, SHOT_DIR / image_name, width_cm)
        return

    anchor = find_paragraph(doc, anchor_text)
    img = add_picture_after(anchor, SHOT_DIR / image_name, width_cm=width_cm)
    cap = add_caption_after(img, caption)
    add_body_after(cap, note)


def rename_caption(doc: Document, old: str, new: str) -> None:
    for para in doc.paragraphs:
        if para.text.strip() == old:
            para.text = new
            run = para.runs[0] if para.runs else para.add_run(new)
            if para.runs:
                for r in para.runs:
                    r.text = ""
                run = para.runs[0]
                run.text = new
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Cm(0)
            set_run_font(run, 10.5, east_asia=HEI_FONT)
            return


def main() -> None:
    doc = Document(DOC_PATH)

    rename_caption(doc, "图5-1 系统功能测试结果统计图", "图5-4 系统功能测试结果统计图")

    insert_figure_block(
        doc,
        "5.2.1 用户认证测试",
        "01-login-page.png",
        "图5-1 用户认证测试过程截图",
        "图5-1展示了用户认证测试过程中的系统登录界面。测试时通过输入账号信息并提交登录请求，验证系统是否能够正确完成身份认证与角色识别。",
    )
    insert_figure_block(
        doc,
        "5.2.2 智能问答测试",
        "05-chat-answer.png",
        "图5-2 智能客服对话测试过程截图",
        "图5-2展示了智能客服对话测试过程。测试中向系统输入企业知识相关问题，系统能够返回模型生成结果，并以对话形式展示问答内容。",
    )
    insert_figure_block(
        doc,
        "5.2.4 后台管理测试",
        "04-admin-users.png",
        "图5-3 后台管理测试过程截图",
        "图5-3展示了后台管理测试过程中的用户管理界面。测试时重点验证用户创建、角色切换、账号启用与禁用以及密码重置等功能是否能够正常执行。",
    )

    doc.save(TARGET_DOC_PATH)
    print(TARGET_DOC_PATH)


if __name__ == "__main__":
    main()
