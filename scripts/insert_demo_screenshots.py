from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DOC_PATH = ROOT / "output" / "doc" / "ibot_thesis_draft.docx"
SHOT_DIR = ROOT / "output" / "playwright"

CN_FONT = "宋体"
HEI_FONT = "黑体"
EN_FONT = "Times New Roman"


def set_run_font(run, size: float, east_asia: str = CN_FONT, ascii_font: str = EN_FONT, bold: bool = False) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def insert_paragraph_after(paragraph: Paragraph, text: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        run = new_para.add_run(text)
        set_run_font(run, 12)
    return new_para


def add_caption_after(paragraph: Paragraph, caption: str) -> Paragraph:
    cap = insert_paragraph_after(paragraph)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(6)
    run = cap.add_run(caption)
    set_run_font(run, 10.5, east_asia=HEI_FONT)
    return cap


def add_body_after(paragraph: Paragraph, text: str) -> Paragraph:
    p = insert_paragraph_after(paragraph)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = Pt(22)
    run = p.add_run(text)
    set_run_font(run, 12)
    return p


def add_picture_after(paragraph: Paragraph, image_path: Path, width_cm: float = 13.2) -> Paragraph:
    p = insert_paragraph_after(paragraph)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    return p


def find_paragraph(doc: Document, text: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip() == text:
            return para
    raise ValueError(f"Could not find paragraph: {text}")


def has_caption(doc: Document, caption: str) -> bool:
    return any(para.text.strip() == caption for para in doc.paragraphs)


def replace_existing_image_paragraph(doc: Document, caption: str, image_path: Path, width_cm: float) -> None:
    paragraphs = list(doc.paragraphs)
    for idx, para in enumerate(paragraphs):
        if para.text.strip() == caption and idx > 0:
            image_para = paragraphs[idx - 1]
            new_p = OxmlElement("w:p")
            image_para._p.addprevious(new_p)
            new_para = Paragraph(new_p, image_para._parent)
            new_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            new_para.paragraph_format.first_line_indent = Cm(0)
            new_para.paragraph_format.left_indent = Cm(0)
            new_para.paragraph_format.space_before = Pt(6)
            new_para.paragraph_format.space_after = Pt(6)
            new_para.paragraph_format.keep_together = True
            new_para.paragraph_format.keep_with_next = True
            new_para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = new_para.add_run()
            run.add_picture(str(image_path), width=Cm(width_cm))
            image_para._element.getparent().remove(image_para._element)
            return


def insert_figure_block(
    doc: Document,
    anchor_text: str,
    image_name: str,
    caption: str,
    note: str,
    width_cm: float = 13.2,
) -> None:
    if has_caption(doc, caption):
        replace_existing_image_paragraph(doc, caption, SHOT_DIR / image_name, width_cm)
        return

    anchor = find_paragraph(doc, anchor_text)
    img_para = add_picture_after(anchor, SHOT_DIR / image_name, width_cm=width_cm)
    cap_para = add_caption_after(img_para, caption)
    add_body_after(cap_para, note)


def main() -> None:
    if not DOC_PATH.exists():
        raise FileNotFoundError(f"Document not found: {DOC_PATH}")

    doc = Document(DOC_PATH)

    insert_figure_block(
        doc,
        "图4-3 知识库上传与重建流程图",
        "03-admin-knowledge.png",
        "图4-4 知识库管理界面",
        "图4-4展示了系统知识库管理界面。管理员可以设置知识库名称、上传知识文档、查看当前检索状态，并在文档列表区域查看已导入的知识内容。",
        width_cm=13.6,
    )
    insert_figure_block(
        doc,
        "管理员可在后台创建用户、切换角色、启用或禁用账户、重置密码并删除账户，从而完成多用户环境下的权限控制。",
        "04-admin-users.png",
        "图4-5 用户管理界面",
        "图4-5展示了系统用户管理界面。管理员能够创建新用户，并对现有用户执行角色切换、账户禁用、密码重置等管理操作。",
        width_cm=13.6,
    )
    insert_figure_block(
        doc,
        "系统设置支持修改默认知识库、检索返回条数、最小相似度阈值和上传限制等参数；任务监控用于展示知识库重建、数据源测试等异步任务的状态和进度。",
        "06-admin-settings.png",
        "图4-6 系统设置界面",
        "图4-6展示了系统设置界面。管理员可以在该页面中维护默认知识库名称、检索参数以及上传大小限制等运行配置。",
        width_cm=13.6,
    )
    insert_figure_block(
        doc,
        "聊天界面负责会话管理、消息展示、输入发送和引用信息渲染。用户可以新建会话、查看历史记录，并在回答区域看到引用的文件片段和相似度分数。",
        "05-chat-answer.png",
        "图4-7 智能客服对话界面",
        "图4-7展示了系统智能客服对话界面。用户可在输入框中提交问题，系统返回模型生成结果，并按会话形式展示对话内容。",
        width_cm=13.6,
    )
    insert_figure_block(
        doc,
        "管理界面以统一的控制台形式组织知识库管理、数据源管理、用户管理、系统设置和任务监控等子页面，便于管理员在同一界面完成系统维护工作。",
        "01-login-page.png",
        "图4-8 系统登录界面",
        "图4-8展示了系统登录界面。系统通过弹窗方式提供登录与注册入口，用户可根据角色选择相应身份后进入聊天工作区或后台管理工作区。",
        width_cm=13.6,
    )

    doc.save(DOC_PATH)
    print(DOC_PATH)


if __name__ == "__main__":
    main()
